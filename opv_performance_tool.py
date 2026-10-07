# -*- coding: utf-8 -*-
"""
OPV器件性能数据表格处理工具
功能：选择数据文件 → 拖拽排序 → 生成格式化Excel表格
"""

import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side


class OPVDataProcessor:
    """OPV数据处理器：解析txt文件并提取性能参数"""
    
    @staticmethod
    def parse_file(filepath):
        """
        解析OPV性能数据文件
        返回: dict {filename, jsc, voc, ff, pce}
        """
        filename = os.path.splitext(os.path.basename(filepath))[0]
        
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                lines = f.readlines()
        except UnicodeDecodeError:
            with open(filepath, 'r', encoding='gbk') as f:
                lines = f.readlines()
        
        # 第1行是列标题，第3行是数据
        if len(lines) < 3:
            raise ValueError(f"文件格式错误: {filepath}")
        
        headers = lines[0].rstrip('\r\n').split('\t')
        values = lines[2].rstrip('\r\n').split('\t')
        
        # 创建列名到值的映射
        data_map = {}
        for i, header in enumerate(headers):
            if i < len(values):
                data_map[header] = values[i]
        
        # 提取数据
        jsc = float(data_map.get('Jsc', 0))
        voc_mv = float(data_map.get('Voc', 0))  # 原始单位: mV
        voc_v = voc_mv / 1000  # 转换为V
        ff = float(data_map.get('FF', 0))
        pce = float(data_map.get('Eta', 0))
        
        return {
            'filename': filename,
            'jsc': round(jsc, 2),
            'voc': round(voc_v, 3),  # Voc保留三位小数
            'ff': round(ff, 2),
            'pce': round(pce, 2)
        }


class DraggableListbox(tk.Listbox):
    """支持拖拽排序的Listbox"""
    
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.bind('<Button-1>', self._on_click)
        self.bind('<B1-Motion>', self._on_drag)
        self.bind('<ButtonRelease-1>', self._on_release)
        self._drag_data = {'index': None}
    
    def _on_click(self, event):
        """记录拖拽起始位置"""
        self._drag_data['index'] = self.nearest(event.y)
        self.selection_clear(0, tk.END)
        self.selection_set(self._drag_data['index'])
    
    def _on_drag(self, event):
        """拖拽过程中更新选中状态"""
        current_index = self.nearest(event.y)
        if current_index != self._drag_data['index'] and self._drag_data['index'] is not None:
            # 获取当前项的值
            item = self.get(self._drag_data['index'])
            # 删除原位置
            self.delete(self._drag_data['index'])
            # 插入新位置
            self.insert(current_index, item)
            # 更新索引
            self._drag_data['index'] = current_index
            # 更新选中状态
            self.selection_clear(0, tk.END)
            self.selection_set(current_index)
    
    def _on_release(self, event):
        """拖拽结束"""
        self._drag_data['index'] = None


class MainApp:
    """主应用程序"""
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("LUEFlow | OPV性能数据表格处理工具")
        self.root.geometry("500x450")
        self.root.resizable(True, True)
        
        self.file_paths = {}  # 文件名 -> 完整路径的映射
        
        self._setup_ui()
        
        # 启动后自动弹出文件选择对话框
        self.root.after(100, self._select_files)
    
    def _setup_ui(self):
        """设置界面"""
        # 标题
        title_label = tk.Label(
            self.root, 
            text="OPV器件性能数据处理", 
            font=("微软雅黑", 14, "bold")
        )
        title_label.pack(pady=10)
        
        # 说明文字
        hint_label = tk.Label(
            self.root,
            text="拖拽文件名调整顺序，顺序将决定表格中的排列",
            font=("微软雅黑", 9),
            fg="gray"
        )
        hint_label.pack()
        
        # 文件列表框架
        list_frame = tk.Frame(self.root)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)
        
        # 滚动条
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 可拖拽列表
        self.listbox = DraggableListbox(
            list_frame,
            yscrollcommand=scrollbar.set,
            font=("Arial", 11),
            selectmode=tk.SINGLE,
            height=12,
            activestyle='none',
            selectbackground='#4a90d9',
            selectforeground='white'
        )
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.listbox.yview)
        
        # 按钮框架
        btn_frame = tk.Frame(self.root)
        btn_frame.pack(pady=10)
        
        # 添加文件按钮
        add_btn = tk.Button(
            btn_frame,
            text="添加文件",
            font=("微软雅黑", 10),
            command=self._select_files,
            width=12
        )
        add_btn.pack(side=tk.LEFT, padx=5)
        
        # 清空列表按钮
        clear_btn = tk.Button(
            btn_frame,
            text="清空列表",
            font=("微软雅黑", 10),
            command=self._clear_list,
            width=12
        )
        clear_btn.pack(side=tk.LEFT, padx=5)
        
        # 生成表格按钮
        generate_btn = tk.Button(
            btn_frame,
            text="生成表格",
            font=("微软雅黑", 10, "bold"),
            command=self._generate_excel,
            width=12,
            bg='#4a90d9',
            fg='white'
        )
        generate_btn.pack(side=tk.LEFT, padx=5)
    
    def _select_files(self):
        """选择文件"""
        filepaths = filedialog.askopenfilenames(
            title="选择OPV性能数据文件",
            filetypes=[("文本文件", "*.txt"), ("所有文件", "*.*")]
        )
        
        if filepaths:
            for fp in filepaths:
                filename = os.path.basename(fp)
                if filename not in self.file_paths:
                    self.file_paths[filename] = fp
                    self.listbox.insert(tk.END, filename)
    
    def _clear_list(self):
        """清空列表"""
        self.listbox.delete(0, tk.END)
        self.file_paths.clear()
    
    def _generate_excel(self):
        """生成Excel表格"""
        if self.listbox.size() == 0:
            messagebox.showwarning("提示", "请先选择数据文件！")
            return
        
        # 按列表顺序获取文件
        ordered_files = [self.listbox.get(i) for i in range(self.listbox.size())]
        
        # 解析所有文件的数据
        data_list = []
        for filename in ordered_files:
            filepath = self.file_paths[filename]
            try:
                data = OPVDataProcessor.parse_file(filepath)
                data_list.append(data)
            except Exception as e:
                messagebox.showerror("错误", f"解析文件失败: {filename}\n{str(e)}")
                return
        
        # 选择保存位置
        save_path = filedialog.asksaveasfilename(
            title="保存表格",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile="OPV性能数据表格.xlsx"
        )
        
        if not save_path:
            return
        
        # 创建Excel
        try:
            self._create_excel(data_list, save_path)
            messagebox.showinfo("成功", f"表格已生成：\n{save_path}")
        except Exception as e:
            messagebox.showerror("错误", f"生成表格失败：\n{str(e)}")
    
    def _create_excel(self, data_list, save_path):
        """创建Excel文件"""
        wb = Workbook()
        ws = wb.active
        ws.title = "OPV性能数据"
        
        # 字体定义
        font_cn = Font(name='微软雅黑', size=11)  # 中文
        font_en = Font(name='Arial', size=11)      # 英文/数字
        font_header_cn = Font(name='微软雅黑', size=11, bold=True)
        font_header_en = Font(name='Arial', size=11, bold=True)
        
        # 对齐方式
        center_align = Alignment(horizontal='center', vertical='center')
        
        # 边框
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 表头
        headers = ['器件名称', 'Jsc (mA/cm²)', 'Voc (V)', 'FF (%)', 'PCE (%)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            # 判断是否为中文表头
            if col == 1:
                cell.font = font_header_cn
            else:
                cell.font = font_header_en
            cell.alignment = center_align
            cell.border = thin_border
        
        # 数据行
        for row_idx, data in enumerate(data_list, 2):
            # 器件名称 (中文字体)
            cell = ws.cell(row=row_idx, column=1, value=data['filename'])
            cell.font = font_cn
            cell.alignment = center_align
            cell.border = thin_border
            
            # Jsc
            cell = ws.cell(row=row_idx, column=2, value=data['jsc'])
            cell.font = font_en
            cell.alignment = center_align
            cell.border = thin_border
            
            # Voc
            cell = ws.cell(row=row_idx, column=3, value=data['voc'])
            cell.font = font_en
            cell.alignment = center_align
            cell.border = thin_border
            
            # FF
            cell = ws.cell(row=row_idx, column=4, value=data['ff'])
            cell.font = font_en
            cell.alignment = center_align
            cell.border = thin_border
            
            # PCE
            cell = ws.cell(row=row_idx, column=5, value=data['pce'])
            cell.font = font_en
            cell.alignment = center_align
            cell.border = thin_border
        
        # 调整列宽
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 12
        
        # 保存
        wb.save(save_path)
    
    def run(self):
        """运行应用"""
        self.root.mainloop()


if __name__ == "__main__":
    app = MainApp()
    app.run()
