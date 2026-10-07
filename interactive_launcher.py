# -*- coding: utf-8 -*-
"""
交互式LUE/AVT计算启动器
功能：启动时选择计算模式（仅AVT 或 LUE+AVT），支持批量文件处理和表格输出
"""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import os
import sys
from pathlib import Path
from difflib import SequenceMatcher

# 确保能够导入lue_calculator模块
sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

from lue_calculator import AVTEngine, LUEEngine
from lue_calculator.data_parser import parse_performance_file


class DraggableListbox(tk.Listbox):
    """支持拖拽排序的Listbox（复用自性能表格处理台）"""
    
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.bind('<Button-1>', self._on_click)
        self.bind('<B1-Motion>', self._on_drag)
        self.bind('<ButtonRelease-1>', self._on_release)
        self._drag_data = {'index': None}
    
    def _on_click(self, event):
        """记录拖拽起始位置"""
        self._drag_data['index'] = self.nearest(event.y)
        self.selection_clear(0, tk.END)
        self.selection_set(self._drag_data['index'])
    
    def _on_drag(self, event):
        """拖拽过程中更新选中状态"""
        current_index = self.nearest(event.y)
        if current_index != self._drag_data['index'] and self._drag_data['index'] is not None:
            item = self.get(self._drag_data['index'])
            self.delete(self._drag_data['index'])
            self.insert(current_index, item)
            self._drag_data['index'] = current_index
            self.selection_clear(0, tk.END)
            self.selection_set(current_index)
    
    def _on_release(self, event):
        """拖拽结束"""
        self._drag_data['index'] = None


class ExcelGenerator:
    """Excel表格生成器（统一样式管理）"""
    
    # 字体定义
    FONT_CN = Font(name='微软雅黑', size=11)
    FONT_EN = Font(name='Arial', size=11)
    FONT_HEADER_CN = Font(name='微软雅黑', size=11, bold=True)
    FONT_HEADER_EN = Font(name='Arial', size=11, bold=True)
    
    # 对齐方式
    CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
    
    # 边框
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    @classmethod
    def apply_cell_style(cls, cell, is_header=False, is_chinese=False):
        """应用单元格样式"""
        if is_header:
            cell.font = cls.FONT_HEADER_CN if is_chinese else cls.FONT_HEADER_EN
        else:
            cell.font = cls.FONT_CN if is_chinese else cls.FONT_EN
        cell.alignment = cls.CENTER_ALIGN
        cell.border = cls.THIN_BORDER
    
    @classmethod
    def generate_avt_table(cls, data_list, save_path):
        """生成AVT表格"""
        wb = Workbook()
        ws = wb.active
        ws.title = "AVT数据"
        
        # 表头
        headers = ['器件名称', 'AVT-Human (%)', 'AVT-Plant (%)']
        header_is_cn = [True, False, False]
        
        for col, (header, is_cn) in enumerate(zip(headers, header_is_cn), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cls.apply_cell_style(cell, is_header=True, is_chinese=is_cn)
        
        # 数据行
        for row_idx, data in enumerate(data_list, 2):
            # 器件名称
            cell = ws.cell(row=row_idx, column=1, value=data['name'])
            cls.apply_cell_style(cell, is_chinese=True)
            
            # AVT-Human
            cell = ws.cell(row=row_idx, column=2, value=round(data['avt_human'], 2))
            cls.apply_cell_style(cell)
            
            # AVT-Plant
            cell = ws.cell(row=row_idx, column=3, value=round(data['avt_plant'], 2))
            cls.apply_cell_style(cell)
        
        # 调整列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        
        wb.save(save_path)
    
    @classmethod
    def generate_lue_table(cls, data_list, save_path):
        """生成LUE+AVT完整表格"""
        wb = Workbook()
        ws = wb.active
        ws.title = "LUE性能数据"
        
        # 表头
        headers = ['器件名称', 'Jsc (mA/cm²)', 'Voc (V)', 'FF (%)', 'PCE (%)',
                   'AVT-Human (%)', 'AVT-Plant (%)', 'LUE-Human (%)', 'LUE-Plant (%)']
        header_is_cn = [True] + [False] * 8
        
        for col, (header, is_cn) in enumerate(zip(headers, header_is_cn), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cls.apply_cell_style(cell, is_header=True, is_chinese=is_cn)
        
        # 数据行
        for row_idx, data in enumerate(data_list, 2):
            values = [
                data['name'],
                round(data['jsc'], 2),
                round(data['voc'], 3),
                round(data['ff'], 2),
                round(data['pce'], 2),
                round(data['avt_human'], 2),
                round(data['avt_plant'], 2),
                round(data['lue_human'], 2),
                round(data['lue_plant'], 2)
            ]
            is_cn_list = [True] + [False] * 8
            
            for col, (value, is_cn) in enumerate(zip(values, is_cn_list), 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cls.apply_cell_style(cell, is_chinese=is_cn)
        
        # 调整列宽
        widths = [20, 15, 12, 10, 12, 15, 15, 15, 15]
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = width
        
        wb.save(save_path)


class AVTModeHandler:
    """仅AVT模式处理器"""
    
    def __init__(self, parent):
        self.parent = parent
        self.window = None
        self.file_paths = {}  # 显示名 -> 完整路径
        self.listbox = None
        self.avt_engine = AVTEngine()
    
    def show(self):
        """显示AVT模式界面"""
        self.window = tk.Toplevel(self.parent)
        self.window.title("AVT计算模式")
        self.window.geometry("550x500")
        self.window.resizable(True, True)
        
        # 将窗口置于父窗口之上
        self.window.transient(self.parent)
        self.window.grab_set()
        
        self._setup_ui()
        
        # 启动后自动弹出文件选择
        self.window.after(100, self._select_files)
    
    def _setup_ui(self):
        """设置界面"""
        # 标题
        title = tk.Label(
            self.window,
            text="📊 AVT计算（平均可见光透过率）",
            font=("微软雅黑", 14, "bold")
        )
        title.pack(pady=15)
        
        # 说明
        hint = tk.Label(
            self.window,
            text="选择透过率(.t)或吸光度(.abs)文件，拖拽调整顺序",
            font=("微软雅黑", 10),
            fg="gray"
        )
        hint.pack()
        
        # 文件列表框架
        list_frame = tk.Frame(self.window)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)
        
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.listbox = DraggableListbox(
            list_frame,
            yscrollcommand=scrollbar.set,
            font=("Arial", 11),
            selectmode=tk.SINGLE,
            height=12,
            activestyle='none',
            selectbackground='#4a90d9',
            selectforeground='white'
        )
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.listbox.yview)
        
        # 按钮框架
        btn_frame = tk.Frame(self.window)
        btn_frame.pack(pady=10)
        
        tk.Button(
            btn_frame, text="添加文件", font=("微软雅黑", 10),
            command=self._select_files, width=12
        ).pack(side=tk.LEFT, padx=5)
        
        tk.Button(
            btn_frame, text="清空列表", font=("微软雅黑", 10),
            command=self._clear_list, width=12
        ).pack(side=tk.LEFT, padx=5)
        
        tk.Button(
            btn_frame, text="生成表格", font=("微软雅黑", 10, "bold"),
            command=self._generate_table, width=12,
            bg='#4a90d9', fg='white'
        ).pack(side=tk.LEFT, padx=5)
    
    def _select_files(self):
        """选择光谱文件"""
        filepaths = filedialog.askopenfilenames(
            title="选择光谱文件（透过率或吸光度）",
            filetypes=[("光谱文件", "*.txt *.TXT"), ("所有文件", "*.*")]
        )
        
        if filepaths:
            for fp in filepaths:
                display_name = os.path.basename(fp)
                if display_name not in self.file_paths:
                    self.file_paths[display_name] = fp
                    self.listbox.insert(tk.END, display_name)
    
    def _clear_list(self):
        """清空列表"""
        self.listbox.delete(0, tk.END)
        self.file_paths.clear()
    
    def _generate_table(self):
        """生成AVT表格"""
        if self.listbox.size() == 0:
            messagebox.showwarning("提示", "请先选择光谱文件！")
            return
        
        # 按顺序获取文件
        ordered_files = [self.listbox.get(i) for i in range(self.listbox.size())]
        
        # 计算AVT
        data_list = []
        for display_name in ordered_files:
            filepath = self.file_paths[display_name]
            try:
                result = self.avt_engine.calculate_both_from_file(filepath)
                # 提取样品名（去掉扩展名和常见后缀）
                name = os.path.splitext(display_name)[0]
                name = name.replace('(UDS)', '').strip()
                
                data_list.append({
                    'name': name,
                    'avt_human': result['human'],
                    'avt_plant': result['plant']
                })
            except Exception as e:
                messagebox.showerror("错误", f"计算失败: {display_name}\n{str(e)}")
                return
        
        # 选择保存位置
        save_path = filedialog.asksaveasfilename(
            title="保存AVT表格",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile="AVT计算结果.xlsx"
        )
        
        if not save_path:
            return
        
        try:
            ExcelGenerator.generate_avt_table(data_list, save_path)
            messagebox.showinfo("成功", f"表格已生成：\n{save_path}")
        except Exception as e:
            messagebox.showerror("错误", f"生成表格失败：\n{str(e)}")


class LUEModeHandler:
    """LUE+AVT模式处理器 - 左右双列布局"""
    
    def __init__(self, parent):
        self.parent = parent
        self.window = None
        self.lue_engine = LUEEngine()
        self.avt_engine = AVTEngine()
        
        # 文件路径映射：显示名 -> 完整路径
        self.spectrum_paths = {}
        self.perf_paths = {}
        
        # 列表框引用
        self.spectrum_listbox = None
        self.perf_listbox = None
    
    def show(self):
        """显示LUE模式界面"""
        self.window = tk.Toplevel(self.parent)
        self.window.title("LUE+AVT计算模式")
        self.window.geometry("800x550")
        self.window.resizable(True, True)
        
        self.window.transient(self.parent)
        self.window.grab_set()
        
        self._setup_ui()
    
    def _setup_ui(self):
        """设置界面"""
        # 标题
        tk.Label(
            self.window,
            text="📈 LUE+AVT完整计算",
            font=("微软雅黑", 14, "bold")
        ).pack(pady=10)
        
        # 说明
        tk.Label(
            self.window,
            text="左侧添加光谱文件，右侧添加性能文件，拖拽调整顺序使行对行配对",
            font=("微软雅黑", 10),
            fg="gray"
        ).pack()
        
        # 主内容区域：左右双列
        main_frame = tk.Frame(self.window)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=15)
        
        # ===== 左侧：光谱文件 =====
        left_frame = tk.Frame(main_frame)
        left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10))
        
        tk.Label(
            left_frame,
            text="📁 光谱文件（透过率/吸光度）",
            font=("微软雅黑", 11, "bold"),
            fg="#3498db"
        ).pack(anchor='w')
        
        # 光谱文件列表
        spectrum_list_frame = tk.Frame(left_frame)
        spectrum_list_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        spectrum_scrollbar = tk.Scrollbar(spectrum_list_frame)
        spectrum_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.spectrum_listbox = DraggableListbox(
            spectrum_list_frame,
            yscrollcommand=spectrum_scrollbar.set,
            font=("Arial", 10),
            selectmode=tk.SINGLE,
            height=15,
            activestyle='none',
            selectbackground='#3498db',
            selectforeground='white'
        )
        self.spectrum_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        spectrum_scrollbar.config(command=self.spectrum_listbox.yview)
        
        # 光谱文件按钮
        spectrum_btn_frame = tk.Frame(left_frame)
        spectrum_btn_frame.pack(fill=tk.X, pady=5)
        
        tk.Button(
            spectrum_btn_frame, text="添加", font=("微软雅黑", 9),
            command=self._add_spectrum_files, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        tk.Button(
            spectrum_btn_frame, text="删除选中", font=("微软雅黑", 9),
            command=self._remove_spectrum_file, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        tk.Button(
            spectrum_btn_frame, text="清空", font=("微软雅黑", 9),
            command=self._clear_spectrum_files, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        self.spectrum_count_label = tk.Label(
            spectrum_btn_frame, text="0 个文件", font=("Arial", 9), fg="gray"
        )
        self.spectrum_count_label.pack(side=tk.RIGHT)
        
        # ===== 中间：配对指示 =====
        middle_frame = tk.Frame(main_frame, width=50)
        middle_frame.pack(side=tk.LEFT, fill=tk.Y, padx=5)
        middle_frame.pack_propagate(False)
        
        tk.Label(
            middle_frame,
            text="↔\n配\n对",
            font=("微软雅黑", 12),
            fg="#95a5a6"
        ).pack(expand=True)
        
        # ===== 右侧：性能文件 =====
        right_frame = tk.Frame(main_frame)
        right_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 0))
        
        tk.Label(
            right_frame,
            text="📊 性能文件（J-V测试）",
            font=("微软雅黑", 11, "bold"),
            fg="#27ae60"
        ).pack(anchor='w')
        
        # 性能文件列表
        perf_list_frame = tk.Frame(right_frame)
        perf_list_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        perf_scrollbar = tk.Scrollbar(perf_list_frame)
        perf_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.perf_listbox = DraggableListbox(
            perf_list_frame,
            yscrollcommand=perf_scrollbar.set,
            font=("Arial", 10),
            selectmode=tk.SINGLE,
            height=15,
            activestyle='none',
            selectbackground='#27ae60',
            selectforeground='white'
        )
        self.perf_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        perf_scrollbar.config(command=self.perf_listbox.yview)
        
        # 性能文件按钮
        perf_btn_frame = tk.Frame(right_frame)
        perf_btn_frame.pack(fill=tk.X, pady=5)
        
        tk.Button(
            perf_btn_frame, text="添加", font=("微软雅黑", 9),
            command=self._add_perf_files, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        tk.Button(
            perf_btn_frame, text="删除选中", font=("微软雅黑", 9),
            command=self._remove_perf_file, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        tk.Button(
            perf_btn_frame, text="清空", font=("微软雅黑", 9),
            command=self._clear_perf_files, width=8
        ).pack(side=tk.LEFT, padx=2)
        
        self.perf_count_label = tk.Label(
            perf_btn_frame, text="0 个文件", font=("Arial", 9), fg="gray"
        )
        self.perf_count_label.pack(side=tk.RIGHT)
        
        # ===== 底部：生成按钮 =====
        bottom_frame = tk.Frame(self.window)
        bottom_frame.pack(fill=tk.X, padx=20, pady=15)
        
        # 配对状态提示
        self.pair_status = tk.Label(
            bottom_frame,
            text="",
            font=("微软雅黑", 10),
            fg="#e74c3c"
        )
        self.pair_status.pack(side=tk.LEFT)
        
        tk.Button(
            bottom_frame,
            text="🚀 生成表格",
            font=("微软雅黑", 12, "bold"),
            command=self._generate_table,
            bg='#27ae60',
            fg='white',
            width=15,
            height=1,
            cursor='hand2'
        ).pack(side=tk.RIGHT)
        
        # 提示信息
        tk.Label(
            self.window,
            text="提示：第1行配第1行，第2行配第2行...请确保两边文件数量一致且顺序正确",
            font=("微软雅黑", 9),
            fg="#7f8c8d"
        ).pack(pady=(0, 10))
    
    def _add_spectrum_files(self):
        """添加光谱文件"""
        filepaths = filedialog.askopenfilenames(
            title="选择光谱文件（透过率或吸光度）",
            filetypes=[("光谱文件", "*.txt *.TXT"), ("所有文件", "*.*")]
        )
        
        if filepaths:
            for fp in filepaths:
                display_name = os.path.basename(fp)
                if display_name not in self.spectrum_paths:
                    self.spectrum_paths[display_name] = fp
                    self.spectrum_listbox.insert(tk.END, display_name)
            
            self._update_counts()
    
    def _remove_spectrum_file(self):
        """删除选中的光谱文件"""
        selection = self.spectrum_listbox.curselection()
        if selection:
            idx = selection[0]
            name = self.spectrum_listbox.get(idx)
            self.spectrum_listbox.delete(idx)
            if name in self.spectrum_paths:
                del self.spectrum_paths[name]
            self._update_counts()
    
    def _clear_spectrum_files(self):
        """清空光谱文件"""
        self.spectrum_listbox.delete(0, tk.END)
        self.spectrum_paths.clear()
        self._update_counts()
    
    def _add_perf_files(self):
        """添加性能文件"""
        filepaths = filedialog.askopenfilenames(
            title="选择器件性能文件",
            filetypes=[("文本文件", "*.txt *.TXT"), ("所有文件", "*.*")]
        )
        
        if filepaths:
            for fp in filepaths:
                display_name = os.path.basename(fp)
                if display_name not in self.perf_paths:
                    self.perf_paths[display_name] = fp
                    self.perf_listbox.insert(tk.END, display_name)
            
            self._update_counts()
    
    def _remove_perf_file(self):
        """删除选中的性能文件"""
        selection = self.perf_listbox.curselection()
        if selection:
            idx = selection[0]
            name = self.perf_listbox.get(idx)
            self.perf_listbox.delete(idx)
            if name in self.perf_paths:
                del self.perf_paths[name]
            self._update_counts()
    
    def _clear_perf_files(self):
        """清空性能文件"""
        self.perf_listbox.delete(0, tk.END)
        self.perf_paths.clear()
        self._update_counts()
    
    def _update_counts(self):
        """更新文件计数和配对状态"""
        spec_count = self.spectrum_listbox.size()
        perf_count = self.perf_listbox.size()
        
        self.spectrum_count_label.config(text=f"{spec_count} 个文件")
        self.perf_count_label.config(text=f"{perf_count} 个文件")
        
        if spec_count == 0 and perf_count == 0:
            self.pair_status.config(text="", fg="#7f8c8d")
        elif spec_count == perf_count:
            self.pair_status.config(text=f"✓ {spec_count} 对文件已配对", fg="#27ae60")
        else:
            diff = abs(spec_count - perf_count)
            self.pair_status.config(
                text=f"⚠ 文件数量不一致（差 {diff} 个）", 
                fg="#e74c3c"
            )
    
    def _generate_table(self):
        """生成表格"""
        spec_count = self.spectrum_listbox.size()
        perf_count = self.perf_listbox.size()
        
        if spec_count == 0:
            messagebox.showwarning("提示", "请先添加光谱文件！")
            return
        
        if perf_count == 0:
            messagebox.showwarning("提示", "请先添加性能文件！")
            return
        
        if spec_count != perf_count:
            result = messagebox.askyesno(
                "文件数量不一致",
                f"光谱文件 {spec_count} 个，性能文件 {perf_count} 个。\n"
                f"将只处理前 {min(spec_count, perf_count)} 对文件。\n\n"
                f"是否继续？"
            )
            if not result:
                return
        
        # 按顺序获取文件配对
        pair_count = min(spec_count, perf_count)
        data_list = []
        
        for i in range(pair_count):
            spec_name = self.spectrum_listbox.get(i)
            perf_name = self.perf_listbox.get(i)
            
            spec_path = self.spectrum_paths[spec_name]
            perf_path = self.perf_paths[perf_name]
            
            # 提取样品名（优先使用光谱文件名）
            display_name = os.path.splitext(spec_name)[0]
            display_name = display_name.replace('(UDS)', '').strip()
            
            try:
                # 解析性能文件
                perf_data = parse_performance_file(perf_path)
                
                # 计算AVT
                avt_result = self.avt_engine.calculate_both_from_file(spec_path)
                
                # 计算LUE
                lue_human = avt_result['human'] * perf_data['pce'] / 100
                lue_plant = avt_result['plant'] * perf_data['pce'] / 100
                
                data_list.append({
                    'name': display_name,
                    'jsc': perf_data['jsc'],
                    'voc': perf_data['voc'],
                    'ff': perf_data['ff'],
                    'pce': perf_data['pce'],
                    'avt_human': avt_result['human'],
                    'avt_plant': avt_result['plant'],
                    'lue_human': lue_human,
                    'lue_plant': lue_plant
                })
            except Exception as e:
                messagebox.showerror(
                    "处理失败", 
                    f"第 {i+1} 对文件处理失败：\n"
                    f"光谱: {spec_name}\n"
                    f"性能: {perf_name}\n\n"
                    f"错误: {str(e)}"
                )
                return
        
        # 选择保存位置
        save_path = filedialog.asksaveasfilename(
            title="保存LUE性能表格",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile="LUE性能数据表格.xlsx"
        )
        
        if not save_path:
            return
        
        try:
            ExcelGenerator.generate_lue_table(data_list, save_path)
            messagebox.showinfo("成功", f"表格已生成：\n{save_path}\n\n共处理 {len(data_list)} 对文件。")
        except Exception as e:
            messagebox.showerror("错误", f"生成表格失败：\n{str(e)}")


class InteractiveLauncher:
    """交互式启动器主界面"""
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("LUEFlow | LUE/AVT 计算系统")
        self.root.geometry("450x350")
        self.root.resizable(False, False)
        
        # 居中显示
        self.root.update_idletasks()
        x = (self.root.winfo_screenwidth() - 450) // 2
        y = (self.root.winfo_screenheight() - 350) // 2
        self.root.geometry(f"450x350+{x}+{y}")
        
        self._setup_ui()
    
    def _setup_ui(self):
        """设置界面"""
        # 标题
        title = tk.Label(
            self.root,
            text="🔬 LUEFlow | LUE/AVT 计算系统",
            font=("微软雅黑", 18, "bold"),
            fg="#2c3e50"
        )
        title.pack(pady=30)
        
        # 副标题
        subtitle = tk.Label(
            self.root,
            text="半透明有机太阳能电池性能分析工具",
            font=("微软雅黑", 10),
            fg="gray"
        )
        subtitle.pack()
        
        # 模式选择框架
        mode_frame = tk.Frame(self.root)
        mode_frame.pack(pady=40)
        
        # AVT模式按钮
        avt_btn = tk.Button(
            mode_frame,
            text="📊 仅计算AVT",
            font=("微软雅黑", 12),
            command=self._start_avt_mode,
            width=18,
            height=2,
            bg='#3498db',
            fg='white',
            activebackground='#2980b9',
            activeforeground='white',
            cursor='hand2'
        )
        avt_btn.pack(pady=10)
        
        # LUE+AVT模式按钮
        lue_btn = tk.Button(
            mode_frame,
            text="📈 计算LUE+AVT",
            font=("微软雅黑", 12),
            command=self._start_lue_mode,
            width=18,
            height=2,
            bg='#27ae60',
            fg='white',
            activebackground='#219a52',
            activeforeground='white',
            cursor='hand2'
        )
        lue_btn.pack(pady=10)
        
        # 版本信息
        version = tk.Label(
            self.root,
            text="LUEFlow v0.1.0 | 2026",
            font=("Arial", 9),
            fg="#bdc3c7"
        )
        version.pack(side=tk.BOTTOM, pady=10)
    
    def _start_avt_mode(self):
        """启动AVT模式"""
        handler = AVTModeHandler(self.root)
        handler.show()
    
    def _start_lue_mode(self):
        """启动LUE模式"""
        handler = LUEModeHandler(self.root)
        handler.show()
    
    def run(self):
        """运行应用"""
        self.root.mainloop()


if __name__ == "__main__":
    app = InteractiveLauncher()
    app.run()
